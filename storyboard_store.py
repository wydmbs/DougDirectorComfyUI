"""
storyboard_store.py — the shared data backbone (see SUITE.md).

Reads/writes a single .xlsx "Assets" registry using the schema generalized
from Pig_and_Rooster_Storyboard_v14.xlsx's Shot List tab, but model-agnostic
and covering both shots and reusable assets (characters, backdrops, props),
keyed by one entry_id convention:
    shot_id      e.g. "1.1", "4.3b"
    CHAR:<name>  e.g. "CHAR:pig"
    MASTER:<name> e.g. "MASTER:farm_field"
    PROP:<name>  e.g. "PROP:crate"

Columns (one row per entry_id; locking again UPDATES the row and appends to
notes rather than duplicating):
    entry_id, entry_type, beat, description,
    prompt_positive, prompt_negative_add,
    model, seed, reused_from, locked, image_path, notes
"""

import os
from datetime import datetime, timezone
from openpyxl import Workbook, load_workbook

COLUMNS = [
    "entry_id", "entry_type", "beat", "description",
    "prompt_positive", "prompt_negative_add",
    "model", "seed", "reused_from", "locked", "image_path", "notes",
]

SHEET_NAME = "Assets"


def _ensure_workbook(path: str):
    if os.path.exists(path):
        wb = load_workbook(path)
        if SHEET_NAME not in wb.sheetnames:
            ws = wb.create_sheet(SHEET_NAME)
            ws.append(COLUMNS)
            wb.save(path)
        return wb
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(COLUMNS)
    wb.save(path)
    return wb


def _col_index(ws, name: str) -> int:
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == name:
            return idx
    raise KeyError(f"Column '{name}' not found in Assets sheet header row.")


def _find_row(ws, entry_id: str, id_col: int):
    for row in range(2, ws.max_row + 1):
        val = ws.cell(row, id_col).value
        if val is not None and str(val) == str(entry_id):
            return row
    return None


def list_entries(path: str) -> list:
    """Return all rows as list of dicts, most-recently-updated note first isn't
    tracked separately here — row order is insertion order, same as the old
    shot list."""
    if not os.path.exists(path):
        return []
    wb = load_workbook(path)
    if SHEET_NAME not in wb.sheetnames:
        return []
    ws = wb[SHEET_NAME]
    rows = []
    for r in range(2, ws.max_row + 1):
        row_vals = [ws.cell(r, c).value for c in range(1, len(COLUMNS) + 1)]
        if all(v is None for v in row_vals):
            continue
        rows.append(dict(zip(COLUMNS, row_vals)))
    return rows


def get_entry(path: str, entry_id: str) -> dict:
    for row in list_entries(path):
        if str(row.get("entry_id")) == str(entry_id):
            return row
    return {}


def lock_entry(
    path: str,
    entry_id: str,
    entry_type: str,
    beat: str,
    description: str,
    prompt_positive: str,
    prompt_negative_add: str,
    model: str,
    seed,
    reused_from: str,
    image_path: str,
    note: str,
) -> None:
    """Write or update one row. A second lock on the same entry_id updates the
    row in place and appends a timestamped line to notes, matching the old
    shot list's inline versioned decision log."""
    wb = _ensure_workbook(path)
    ws = wb[SHEET_NAME]
    id_col = _col_index(ws, "entry_id")
    row = _find_row(ws, entry_id, id_col)

    ts = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    note_line = f"[{ts}] {note}" if note else f"[{ts}] locked"

    if row is None:
        row = ws.max_row + 1
        ws.cell(row, _col_index(ws, "entry_id"), entry_id)
        existing_notes = ""
    else:
        existing_notes = ws.cell(row, _col_index(ws, "notes")).value or ""

    values = {
        "entry_type": entry_type,
        "beat": beat,
        "description": description,
        "prompt_positive": prompt_positive,
        "prompt_negative_add": prompt_negative_add,
        "model": model,
        "seed": seed,
        "reused_from": reused_from,
        "locked": True,
        "image_path": image_path,
        "notes": (existing_notes + "\n" + note_line).strip() if existing_notes else note_line,
    }
    for col_name, val in values.items():
        ws.cell(row, _col_index(ws, col_name), val)

    wb.save(path)
