"""Harry the Advisor: source analysis, provider calls, and editable Build drafts."""

import json
import os
import re
import shutil
import subprocess
import urllib.error
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

import yaml

PROJECT_ROOT = Path(__file__).resolve().parent
LIBRARY_DIR = PROJECT_ROOT / "harry_library"


def set_library_dir(path):
    global LIBRARY_DIR
    LIBRARY_DIR = Path(path)


def _azure_settings():
    """Azure settings come from this app's own config, edited in Setup.

    Environment variables still win when present, so a machine that already
    exports them doesn't need the config file filled in at all.
    """
    import config as cfgmod

    cfg = cfgmod.load_config().azure
    endpoint = os.environ.get("AZURE_OPENAI_ENDPOINT") or cfg.endpoint
    deployment = os.environ.get("AZURE_OPENAI_DEPLOYMENT") or cfg.deployment
    api_version = os.environ.get("AZURE_OPENAI_API_VERSION") or cfg.api_version
    key = os.environ.get(cfg.api_key_env) or os.environ.get("AZURE_OPENAI_API_KEY")
    return {
        "endpoint": (endpoint or "").rstrip("/"),
        "deployment": deployment or "",
        "api_version": api_version or "2024-10-21",
        "key": key or "",
        "key_env": cfg.api_key_env,
    }

HARRY_PERSONA = """You are Harry the Advisor, a thoughtful pre-production advisor for AI-assisted film.
Read the supplied script, story, poem, or narration transcript."""

ERA_NOTE = """
IF AN ERA/SETTING IS SPECIFIED, TREAT IT AS A HARD CONSTRAINT:
Ground every visual detail, object, and technology in that era. Do not introduce anachronistic
equipment (e.g. a motor vehicle in a pre-automobile setting, electric lighting in a candlelit era).
When inferring something the text doesn't name outright, the era is exactly what should decide
which real-world version of it is correct -- if animals are "moved from field to ship" in an era
before motor transport, the era tells you that's a horse and cart, not a truck."""

ITEM_SHAPE_NOTE = """Each item must have this exact shape:
{
  "kind": "%s",
  "name": "human readable name",
  "suggested_id": "%s",
  "description": "what it is and why it matters",
  "continuity_note": "visual or narrative continuity guidance",
  "positive_prompt": "initial target-model-aware visual prompt",
  "negative_prompt": "things to avoid",
  "beat": "the story section this is introduced or most associated with",
  "reused_from": "optional prior item id",
  "camera_direction": "for SHOT only: concrete framing, movement, and focal action; empty otherwise",
  "lighting_direction": "for SHOT only: time, source, contrast, and mood; empty otherwise"
}"""

CHARACTER_SYSTEM_PROMPT = f"""{HARRY_PERSONA}

TASK: identify ONLY the characters in the material -- do not identify backdrops, props, or shots in this pass; those come in separate passes. Do not invent named people without textual evidence.
{ERA_NOTE}

Return JSON only, with this exact shape:
{{"items": [ {ITEM_SHAPE_NOTE % ('CHARACTER', 'CHARACTER:snake_case')} ]}}"""

BACKDROP_SYSTEM_PROMPT = f"""{HARRY_PERSONA}

TASK: identify ONLY recurring backdrops/locations in the material (kind "BACKDROP") -- do not identify characters, props, or shots in this pass; those are handled separately. You will be told which characters were already identified in an earlier pass, for context only.
{ERA_NOTE}

A location the story clearly requires but never directly describes (e.g. a farmhouse implied by "farrow & field") MAY be suggested as a tentative item, but mark its continuity_note as "inferred, not explicit in source -- confirm before locking" so the director knows it's a suggestion, not something read directly off the page.

Return JSON only, with this exact shape:
{{"items": [ {ITEM_SHAPE_NOTE % ('BACKDROP', 'BACKDROP:snake_case')} ]}}"""

PROP_SYSTEM_PROMPT = f"""{HARRY_PERSONA}

TASK: identify ONLY recurring key objects/props in the material (kind "PROP") -- do not identify characters, backdrops, or shots in this pass; those are handled separately. You will be told which characters and backdrops were already identified in earlier passes, for context only.
{ERA_NOTE}

RECURRING PHYSICAL OBJECTS ARE PROPS, EVEN WHEN DESCRIBED POETICALLY:
A recurring object central to multiple beats (a vehicle, a vessel, a key object the characters interact with repeatedly) warrants its own PROP entry even if the text never uses a single plain noun for it consistently -- reading "sail ship," "cargo hold," "the hull," and "a sleeping hulk" across a few lines is the same object described four ways, not four separate ideas. Do not require one literal name to appear before recommending something the narrative obviously depends on.
Separately, when the text implies a real-world object or mechanism the story logically requires but never actually states (e.g. characters are "moved from field to ship" with no vehicle named), you MAY suggest it as a tentative item, but mark its continuity_note as "inferred, not explicit in source -- confirm before locking" so the director knows it's a suggestion, not something read directly off the page.

Return JSON only, with this exact shape:
{{"items": [ {ITEM_SHAPE_NOTE % ('PROP', 'PROP:snake_case')} ]}}"""

SHOT_SYSTEM_PROMPT = f"""{HARRY_PERSONA}

TASK: break the material into beats/scenes and produce shot drafts (kind "SHOT") -- one pass, focused entirely on scene coverage, since this is the task most likely to get shortchanged when asked alongside everything else. You will be told which characters, backdrops, and props were already identified in earlier passes -- reference them by name in shot descriptions/continuity notes for continuity, don't redefine them.
{ERA_NOTE}

SCENE/BEAT COVERAGE IS MANDATORY, NOT OPTIONAL:
If the source text contains explicit section markers (bracketed headers like "< The Ship >", chapter titles, scene headings, or similarly clear structural breaks), treat each one as a beat and include at least one SHOT item for every single marked section -- do not silently skip a marked section because it seems minor or because you are trying to keep the list short. A missing beat is a more serious omission than a slightly longer list. If the source has no explicit markers, infer a reasonable beat structure yourself (opening, rising action, climax, resolution, etc.) and still cover it with SHOT items.

Return JSON only, with this exact shape:
{{
  "summary": "short production reading of the whole piece",
  "questions": ["only essential clarification questions; empty when the material is clear"],
  "items": [ {ITEM_SHAPE_NOTE % ('SHOT', 'beat.shot, e.g. 1.1')} ]
}}"""


class HarryError(Exception):
    pass


class HarryRateLimitError(HarryError):
    """The provider is throttling; retrying another analysis pass worsens it."""

    def __init__(self, message, retry_after=""):
        super().__init__(message)
        self.retry_after = retry_after


def _timestamp():
    return datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")


def _slug(value):
    return "_".join(re.sub(r"[^a-z0-9]+", " ", (value or "").lower()).split())


def provider_status(provider):
    provider = provider or "Claude"
    if provider == "Azure OpenAI":
        azure = _azure_settings()
        if azure["endpoint"] and azure["deployment"] and azure["key"]:
            return "Azure OpenAI is ready."
        missing = []
        if not azure["endpoint"]:
            missing.append("endpoint")
        if not azure["deployment"]:
            missing.append("deployment")
        if not azure["key"]:
            missing.append(f"the {azure['key_env']} environment variable")
        return "Azure OpenAI still needs " + ", ".join(missing) + " (set these in Setup)."
    env_map = {"Claude": "ANTHROPIC_API_KEY", "OpenAI": "OPENAI_API_KEY", "Grok": "XAI_API_KEY"}
    if provider == "Ollama":
        return "Ollama runs locally at its configured address; no material leaves this machine."
    return f"{provider} is ready." if os.environ.get(env_map[provider]) else f"Set {env_map[provider]} in the app process environment to use {provider}."


def _request(url, headers, payload, timeout=120):
    request = urllib.request.Request(url, data=json.dumps(payload).encode("utf-8"), headers=headers)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return json.loads(response.read().decode("utf-8"))
    except urllib.error.HTTPError as error:
        detail = error.read().decode("utf-8", errors="replace")[:600]
        retry_after = error.headers.get("Retry-After", "") if error.headers else ""
        if error.code == 429:
            wait = f" Retry after {retry_after} seconds." if retry_after else " Wait briefly, then try again."
            raise HarryRateLimitError(
                "The selected model is rate limited by its provider." + wait,
                retry_after=retry_after,
            ) from error
        raise HarryError(f"Provider returned HTTP {error.code}: {detail}") from error
    except urllib.error.URLError as error:
        raise HarryError(f"Could not reach the provider: {error.reason}") from error


def _chat(provider, system_prompt, user_prompt):
    if provider == "Azure OpenAI":
        azure = _azure_settings()
        key = azure["key"]
        if not key:
            raise HarryError("Azure API key is not available to this app process. Set %s and restart the app." % azure["key_env"])
        if not azure["endpoint"] or not azure["deployment"]:
            raise HarryError("Azure endpoint and deployment are not configured yet -- fill them in on the Setup tab.")
        endpoint = azure["endpoint"]
        url = f"{endpoint}/openai/v1/chat/completions"
        response = _request(url, {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}, {
            "model": azure["deployment"], "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}], "max_completion_tokens": 5000,
        })
        return response["choices"][0]["message"]["content"]
    if provider == "Claude":
        key = os.environ.get("ANTHROPIC_API_KEY")
        if not key:
            raise HarryError("ANTHROPIC_API_KEY is not available to this app process.")
        response = _request("https://api.anthropic.com/v1/messages", {"x-api-key": key, "anthropic-version": "2023-06-01", "content-type": "application/json"}, {
            "model": os.environ.get("HARRY_CLAUDE_MODEL", "claude-sonnet-4-5"), "max_tokens": 5000,
            "system": system_prompt, "messages": [{"role": "user", "content": user_prompt}],
        })
        return "".join(block.get("text", "") for block in response.get("content", []) if block.get("type") == "text")
    if provider == "OpenAI":
        key = os.environ.get("OPENAI_API_KEY")
        if not key:
            raise HarryError("OPENAI_API_KEY is not available to this app process.")
        response = _request("https://api.openai.com/v1/chat/completions", {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}, {
            "model": os.environ.get("HARRY_OPENAI_MODEL", "gpt-4.1"), "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}], "response_format": {"type": "json_object"},
        })
        return response["choices"][0]["message"]["content"]
    if provider == "Grok":
        key = os.environ.get("XAI_API_KEY")
        if not key:
            raise HarryError("XAI_API_KEY is not available to this app process.")
        response = _request("https://api.x.ai/v1/chat/completions", {"Authorization": f"Bearer {key}", "Content-Type": "application/json"}, {
            "model": os.environ.get("HARRY_GROK_MODEL", "grok-3"), "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}], "response_format": {"type": "json_object"},
        })
        return response["choices"][0]["message"]["content"]
    response = _request(os.environ.get("HARRY_OLLAMA_URL", "http://127.0.0.1:11434") + "/api/chat", {"Content-Type": "application/json"}, {
        "model": os.environ.get("HARRY_OLLAMA_MODEL", "qwen2.5:7b"), "stream": False,
        "messages": [{"role": "system", "content": system_prompt}, {"role": "user", "content": user_prompt}], "format": "json",
    })
    return response["message"]["content"]


def _parse_json_block(raw):
    fenced = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw.strip(), flags=re.I)
    try:
        return json.loads(fenced)
    except json.JSONDecodeError as error:
        raise HarryError("Harry returned a response that was not valid JSON. Try again or choose another provider.") from error


def _clean_items(raw_items):
    clean = []
    for item in raw_items or []:
        if not isinstance(item, dict) or item.get("kind") not in {"CHARACTER", "BACKDROP", "PROP", "SHOT"}:
            continue
        name = str(item.get("name") or "").strip()
        if not name:
            continue
        kind = item["kind"]
        prefix = {"CHARACTER": "CHARACTER", "BACKDROP": "BACKDROP", "PROP": "PROP"}.get(kind)
        suggested = str(item.get("suggested_id") or "").strip()
        if prefix and not suggested.startswith(prefix + ":"):
            suggested = f"{prefix}:{_slug(name)}"
        if kind == "SHOT" and not suggested:
            suggested = _slug(name) or "draft_shot"
        item["suggested_id"] = suggested
        clean.append({key: str(item.get(key) or "") for key in ("kind", "name", "suggested_id", "description", "continuity_note", "positive_prompt", "negative_prompt", "beat", "reused_from", "camera_direction", "lighting_direction")})
    return clean


def _parse_plan(raw):
    plan = _parse_json_block(raw)
    if not isinstance(plan, dict) or not isinstance(plan.get("items"), list):
        raise HarryError("Harry's response did not contain a usable recommendation list.")
    plan["items"] = _clean_items(plan["items"])
    plan["summary"] = str(plan.get("summary") or "")
    plan["questions"] = [str(question) for question in plan.get("questions", []) if str(question).strip()]
    return plan


def _parse_items(raw):
    """Same cleaning as _parse_plan, but for a stage call that returns only
    {"items": [...]} with no summary/questions expected."""
    data = _parse_json_block(raw)
    if not isinstance(data, dict) or not isinstance(data.get("items"), list):
        raise HarryError("Harry's response did not contain a usable item list for this pass.")
    return _clean_items(data["items"])


def save_source(title, text, uploaded_path=None):
    LIBRARY_DIR.mkdir(exist_ok=True)
    source_id = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{_slug(title) or 'source'}"
    folder = LIBRARY_DIR / source_id
    folder.mkdir()
    text_path = folder / "source.txt"
    text_path.write_text(text or "", encoding="utf-8")
    if uploaded_path:
        shutil.copy2(uploaded_path, folder / Path(uploaded_path).name)
    metadata = {"source_id": source_id, "title": title or source_id, "created_at": _timestamp(), "text_path": str(text_path), "audio_path": str(folder / Path(uploaded_path).name) if uploaded_path else ""}
    (folder / "metadata.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    return metadata


def extract_text_document(file_path):
    if not file_path:
        raise HarryError("Attach a .txt, .md, .docx, or .pdf document first.")
    path = Path(file_path)
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md", ".rtf"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        try:
            from docx import Document
            from docx.oxml.ns import qn
            from docx.table import Table
            from docx.text.paragraph import Paragraph
        except ImportError as error:
            raise HarryError("Reading .docx needs python-docx. Install it with: pip install python-docx") from error
        doc = Document(path)
        parts = []
        # Walk the document body in actual document order, handling both
        # paragraphs AND tables. document.paragraphs alone only sees
        # top-level body paragraphs -- text laid out in a table (a common
        # choice for scripts/treatments, since it keeps character names
        # and dialogue aligned) is otherwise silently invisible and
        # extraction returns an empty string with no error at all.
        for child in doc.element.body.iterchildren():
            if child.tag == qn("w:p"):
                para = Paragraph(child, doc)
                if para.text.strip():
                    parts.append(para.text)
            elif child.tag == qn("w:tbl"):
                table = Table(child, doc)
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        return "\n".join(parts).strip()
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as error:
            raise HarryError("Reading PDFs needs pypdf. Install it with: pip install pypdf") from error
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages).strip()
    raise HarryError("Supported text documents are .txt, .md, .docx, and .pdf.")


def transcribe_audio(audio_path):
    if not audio_path:
        raise HarryError("Upload audio before asking for transcription.")
    try:
        import whisper
    except ImportError as error:
        raise HarryError("Local transcription needs openai-whisper. Install it with: pip install openai-whisper") from error
    model = whisper.load_model(os.environ.get("HARRY_WHISPER_MODEL", "medium.en"))
    result = model.transcribe(audio_path, language="en")
    return result.get("text", "").strip()


def label_script_passages(provider, passages):
    """Give structurally split source passages meaningful, script-aware labels."""
    if not passages:
        return []
    payload = [{"index": index + 1, "text": excerpt[:900]} for index, (_, excerpt) in enumerate(passages[:24])]
    system = """You are Harry The Helper, a film assistant director.
Give every script passage a concise, meaningful production label. The label must name the dramatic action, setting, or story turn - not repeat the first words and not use generic labels such as Passage 1. Use 2-7 words. Preserve the exact index. Return JSON only: {"passages":[{"index":1,"label":"..."}]}."""
    try:
        data = _parse_json_block(_chat(provider, system, json.dumps(payload)))
    except HarryError:
        return [{"label": label, "excerpt": excerpt} for label, excerpt in passages]
    labels = {int(item.get("index")): str(item.get("label") or "").strip() for item in data.get("passages", []) if str(item.get("index", "")).isdigit()}
    return [{"label": labels.get(index + 1) or label, "excerpt": excerpt} for index, (label, excerpt) in enumerate(passages)]


def playbook_path():
    return LIBRARY_DIR / "directors_playbook.json"


def load_playbook():
    path = playbook_path()
    if not path.exists():
        return []
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        return [item for item in data.get("rules", []) if isinstance(item, dict) and item.get("rule")]
    except (json.JSONDecodeError, OSError):
        return []


def playbook_context():
    rules = load_playbook()
    if not rules:
        return ""
    lines = ["DIRECTOR PLAYBOOK - learned standing decisions. Apply these unless the current script directly conflicts:"]
    lines.extend("- " + item["rule"] for item in rules[-12:])
    return "\n".join(lines) + "\n\n"


def learn_from_feedback(provider, plan, feedback, source_excerpt=""):
    system = """You are Harry The Helper maintaining a director's production playbook.
Extract ONE concise reusable creative or production rule from the feedback and amended call sheet. The rule must help a future script breakdown make a better proposal before the director repeats the note. Do not record one-off plot facts. Return JSON only: {"rule": "...", "reason": "..."}."""
    user = "DIRECTOR FEEDBACK:\n" + (feedback or "") + "\n\nSCRIPT CONTEXT:\n" + (source_excerpt or "") + "\n\nAMENDED CALL SHEET:\n" + json.dumps(plan or {}, indent=2)
    try:
        data = _parse_json_block(_chat(provider, system, user))
    except HarryError:
        return None
    rule = str(data.get("rule") or "").strip()
    reason = str(data.get("reason") or "").strip()
    if not rule:
        return None
    rules = load_playbook()
    normalized = " ".join(rule.lower().split())
    if any(" ".join(str(item.get("rule", "")).lower().split()) == normalized for item in rules):
        return None
    rules.append({"rule": rule, "reason": reason, "learned_at": _timestamp()})
    LIBRARY_DIR.mkdir(exist_ok=True)
    playbook_path().write_text(json.dumps({"rules": rules}, indent=2), encoding="utf-8")
    return rule


def _context_line(label, items):
    if not items:
        return f"{label}: none identified yet."
    names = ", ".join(f"{item['name']} ({item['suggested_id']})" for item in items)
    return f"{label}: {names}"


def _write_analysis_diagnostic(stage_name, error, raw=""):
    """Persist enough local evidence to diagnose provider failures without leaking it to UI."""
    try:
        folder = LIBRARY_DIR / "diagnostics"
        folder.mkdir(parents=True, exist_ok=True)
        path = folder / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{_slug(stage_name)}.json"
        path.write_text(json.dumps({"stage": stage_name, "error": str(error), "raw_tail": (raw or "")[-4000:]}, indent=2), encoding="utf-8")
    except OSError:
        pass


def _run_stage(provider, system_prompt, user_prompt, stage_name, warnings, on_progress=None):
    """Run one structured pass, recovering once from malformed provider JSON."""
    raw = ""
    try:
        if on_progress:
            on_progress(f"Harry The Helper is reviewing {stage_name.lower()}.")
        raw = _chat(provider, system_prompt, user_prompt)
        items = _parse_items(raw)
        if on_progress:
            on_progress(f"{stage_name} complete: {len(items)} proposal(s) found.")
        return items
    except HarryRateLimitError:
        if on_progress:
            on_progress(f"{stage_name} paused: the provider is rate limiting requests. No retry was sent.")
        raise
    except HarryError as first_error:
        if on_progress:
            on_progress(f"{stage_name} response was malformed. Harry The Helper is retrying in a compact format.")
        recovery = user_prompt + "\n\nRECOVERY FORMAT: Return valid JSON only. Keep every field concise. No markdown or explanation outside the JSON object."
        try:
            raw = _chat(provider, system_prompt, recovery)
            items = _parse_items(raw)
            if on_progress:
                on_progress(f"{stage_name} recovered: {len(items)} proposal(s) found.")
            return items
        except HarryRateLimitError:
            if on_progress:
                on_progress(f"{stage_name} paused: the provider began rate limiting during recovery. No further retry was sent.")
            raise
        except HarryError as retry_error:
            _write_analysis_diagnostic(stage_name, retry_error, raw)
            warnings.append(f"{stage_name} pass failed after retry: {retry_error}")
            if on_progress:
                on_progress(f"{stage_name} could not be completed after retry.")
            return []


def _run_shot_stage(provider, user_prompt, on_progress=None):
    """Coverage is mandatory, so retry once in a compact shape before failing.

    Shot cards carry more fields than asset proposals. A long, detailed source
    can therefore exceed a provider's response budget and leave a truncated
    JSON object. The retry keeps the same planning job but asks for concise
    field values and a bounded number of shots, which is more reliable than
    accepting a call sheet with no coverage.
    """
    raw = ""
    try:
        if on_progress:
            on_progress("Harry The Helper is planning camera and lighting coverage.")
        raw = _chat(provider, SHOT_SYSTEM_PROMPT, user_prompt)
        shots = _parse_plan(raw)
    except HarryRateLimitError:
        if on_progress:
            on_progress("Coverage paused: the provider is rate limiting requests. No compact retry was sent.")
        raise
    except HarryError as first_error:
        compact_prompt = user_prompt + """

RECOVERY FORMAT: Your previous response could not be parsed. Return valid JSON only.
Create 8-12 essential SHOT items, or one for every explicit marked section when there are fewer than 8. Keep description, continuity_note, positive_prompt, and negative_prompt concise (one sentence each). camera_direction is required for every shot. Do not add markdown, explanation, or text outside the JSON object."""
        try:
            if on_progress:
                on_progress("Coverage response was malformed. Harry The Helper is retrying a compact shot plan.")
            raw = _chat(provider, SHOT_SYSTEM_PROMPT, compact_prompt)
            shots = _parse_plan(raw)
        except HarryRateLimitError:
            if on_progress:
                on_progress("Coverage paused: the provider is rate limiting requests during recovery. No further retry was sent.")
            raise
        except HarryError as retry_error:
            _write_analysis_diagnostic("Coverage", retry_error, raw)
            raise HarryError(
                "Coverage planning failed twice, so Harry The Helper refused to save an incomplete call sheet. "
                f"First attempt: {first_error}. Retry: {retry_error}"
            ) from retry_error
    missing_camera = [item.get("name", "unnamed shot") for item in shots["items"] if not item.get("camera_direction", "").strip()]
    missing_lighting = [item.get("name", "unnamed shot") for item in shots["items"] if not item.get("lighting_direction", "").strip()]
    if not shots["items"]:
        raise HarryError("Coverage planning returned no shots. Harry The Helper refused to save an incomplete call sheet.")
    if missing_camera or missing_lighting:
        missing = []
        if missing_camera:
            missing.append("camera direction: " + ", ".join(missing_camera[:4]))
        if missing_lighting:
            missing.append("lighting direction: " + ", ".join(missing_lighting[:4]))
        raise HarryError(
            "Coverage planning returned incomplete shot direction (" + "; ".join(missing) + "). "
            "Harry The Helper refused to save an incomplete call sheet."
        )
    if on_progress:
        on_progress(f"Coverage complete: {len(shots['items'])} camera-and-lighting shot proposal(s) found.")
    return shots


def resume_coverage(provider, checkpoint, on_progress=None):
    """Finish a rate-limited brief from its saved cast/world/prop checkpoint."""
    checkpoint = checkpoint or {}
    source_text = str(checkpoint.get("source_text") or "")
    if not source_text.strip():
        raise HarryError("The saved coverage checkpoint has no source text.")
    characters = checkpoint.get("characters") or []
    backdrops = checkpoint.get("backdrops") or []
    props = checkpoint.get("props") or []
    if not (characters or backdrops or props):
        raise HarryError("The saved coverage checkpoint has no completed production passes.")
    title = checkpoint.get("title") or "Untitled project"
    era = checkpoint.get("era") or ""
    notice = "The source stays local." if provider == "Ollama" else f"The source text will be sent to {provider} for this analysis."
    era_line = f"ERA / SETTING: {era}\n" if era else ""
    base = f"PROJECT TITLE: {title}\n{era_line}\n{playbook_context()}SOURCE:\n{source_text}\n\n{notice}\n"
    shot_prompt = (
        base + _context_line("Characters", characters) + "\n"
        + _context_line("Backdrops", backdrops) + "\n"
        + _context_line("Props", props) + "\nReturn the JSON now."
    )
    if on_progress:
        on_progress("Resuming only camera and lighting coverage from the saved cast, world, and props checkpoint.")
    shot_plan = _run_shot_stage(provider, shot_prompt, on_progress)
    items = characters + backdrops + props + shot_plan["items"]
    summary = shot_plan.get("summary") or (
        f"Resumed coverage with {len(characters)} character(s), {len(backdrops)} backdrop(s), "
        f"{len(props)} prop(s), and {len(shot_plan['items'])} shot(s)."
    )
    plan = {"summary": summary, "questions": shot_plan.get("questions") or [], "items": items}
    plan.update({"plan_id": datetime.now().strftime("%Y%m%d_%H%M%S"), "title": title, "provider": provider, "era": era, "created_at": _timestamp(), "source_path": checkpoint.get("source_path", "")})
    LIBRARY_DIR.mkdir(exist_ok=True)
    (LIBRARY_DIR / f"plan_{plan['plan_id']}.json").write_text(json.dumps(plan, indent=2), encoding="utf-8")
    return plan


def analyze_staged(provider, title, source_text, source_path="", uploaded_path=None, era="", on_progress=None):
    """Yield Harry's four real research passes and finally the completed plan."""
    if not source_text.strip():
        raise HarryError("Paste script/story text, or transcribe the uploaded audio first.")
    if not source_path:
        source_path = save_source(title, source_text, uploaded_path).get("source_id", "")
    notice = "The source stays local." if provider == "Ollama" else f"The source text will be sent to {provider} for this analysis."
    era_line = f"ERA / SETTING: {era}\n" if (era or "").strip() else ""
    base = f"PROJECT TITLE: {title or 'Untitled project'}\n{era_line}\n{playbook_context()}SOURCE:\n{source_text}\n\n{notice}\n"
    warnings = []

    character_items = _run_stage(provider, CHARACTER_SYSTEM_PROMPT, base + "Return the JSON now.", "Characters", warnings, on_progress)
    yield {"stage": "cast", "items": character_items, "source_path": source_path}

    backdrop_items = _run_stage(
        provider, BACKDROP_SYSTEM_PROMPT,
        base + _context_line("Already-identified characters", character_items) + "\nReturn the JSON now.",
        "Backdrops", warnings, on_progress,
    )
    yield {"stage": "world", "items": backdrop_items, "source_path": source_path}

    prop_items = _run_stage(
        provider, PROP_SYSTEM_PROMPT,
        base + _context_line("Already-identified characters", character_items) + "\n"
        + _context_line("Already-identified backdrops", backdrop_items) + "\nReturn the JSON now.",
        "Props", warnings, on_progress,
    )
    yield {"stage": "props", "items": prop_items, "source_path": source_path}

    shot_user_prompt = (
        base + _context_line("Characters", character_items) + "\n"
        + _context_line("Backdrops", backdrop_items) + "\n"
        + _context_line("Props", prop_items) + "\nReturn the JSON now."
    )
    shot_plan = _run_shot_stage(provider, shot_user_prompt, on_progress)
    shot_items = shot_plan["items"]
    summary = shot_plan["summary"]
    questions = shot_plan["questions"]

    all_items = character_items + backdrop_items + prop_items + shot_items
    if not all_items:
        raise HarryError("Harry could not produce any recommendations. " + " ".join(warnings))
    if not summary:
        summary = (f"Identified {len(character_items)} character(s), {len(backdrop_items)} backdrop(s), "
                   f"{len(prop_items)} prop(s), and {len(shot_items)} shot(s) across separate focused passes.")
    plan = {"summary": summary, "questions": questions + warnings, "items": all_items}
    plan.update({"plan_id": datetime.now().strftime("%Y%m%d_%H%M%S"), "title": title or "Untitled project", "provider": provider, "era": era or "", "created_at": _timestamp(), "source_path": source_path})
    LIBRARY_DIR.mkdir(exist_ok=True)
    (LIBRARY_DIR / f"plan_{plan['plan_id']}.json").write_text(json.dumps(plan, indent=2), encoding="utf-8")
    yield {"stage": "complete", "plan": plan, "source_path": source_path}


def analyze(provider, title, source_text, source_path="", uploaded_path=None, era="", on_progress=None):
    plan = None
    for event in analyze_staged(provider, title, source_text, source_path, uploaded_path, era, on_progress):
        if event["stage"] == "complete":
            plan = event["plan"]
    if plan is None:
        raise HarryError("Harry could not produce a call sheet.")
    return plan


def _validate_shot_directions(items):
    missing = []
    for item in items:
        if item.get("kind") != "SHOT":
            continue
        absent = [name for name in ("camera_direction", "lighting_direction") if not str(item.get(name) or "").strip()]
        if absent:
            missing.append(item.get("name", "unnamed shot") + " (" + ", ".join(absent) + ")")
    if missing:
        raise HarryError("The amendment returned incomplete coverage direction: " + ", ".join(missing[:4]))


def _proposal_order(item):
    kind_order = {"CHARACTER": 0, "BACKDROP": 1, "PROP": 2, "SHOT": 3}
    suggested_id = str(item.get("suggested_id") or "")
    if item.get("kind") != "SHOT":
        return kind_order.get(item.get("kind"), 9), 0, (), suggested_id.lower()
    parts = []
    for token in re.findall(r"\d+|[a-z]+", suggested_id.lower()):
        parts.append((0, int(token)) if token.isdigit() else (1, token))
    return kind_order["SHOT"], 0, tuple(parts), suggested_id.lower()


def _merge_revision(plan, updates, additions, summary, questions):
    existing = [dict(item) for item in (plan or {}).get("items", [])]
    by_id = {item.get("suggested_id"): index for index, item in enumerate(existing) if item.get("suggested_id")}
    for item in updates:
        key = item.get("suggested_id")
        if key not in by_id:
            raise HarryError("Harry tried to update an unknown proposal: " + str(key))
        current = existing[by_id[key]]
        existing[by_id[key]] = {**current, **{field: value for field, value in item.items() if value != ""}}
    for item in additions:
        key = item.get("suggested_id")
        if key and key in by_id:
            raise HarryError("Harry returned a duplicate new proposal: " + key)
        existing.append(item)
        if key:
            by_id[key] = len(existing) - 1
    existing.sort(key=_proposal_order)
    _validate_shot_directions(existing)
    revised = dict(plan or {})
    revised["items"] = existing
    revised["summary"] = str(summary or plan.get("summary") or "")
    revised["questions"] = [str(question) for question in (questions or []) if str(question).strip()]
    return revised


def _terms(value):
    return {word for word in re.findall(r"[a-z0-9]{3,}", (value or "").lower()) if word not in {"about", "after", "against", "and", "are", "but", "for", "from", "into", "its", "not", "only", "our", "that", "the", "this", "they", "with", "would"}}


def _amendment_context(plan, feedback, source_excerpt, selected_ids=None):
    """Return only the live records needed for one grounded amendment."""
    items = [dict(item) for item in (plan or {}).get("items", [])]
    terms = _terms(feedback) | _terms(source_excerpt)
    scored = []
    for index, item in enumerate(items):
        text = " ".join(str(item.get(field) or "") for field in ("kind", "name", "suggested_id", "description", "continuity_note", "beat", "reused_from"))
        score = len(terms & _terms(text))
        if score:
            scored.append((score, index, item))
    selected = {value for value in (selected_ids or []) if value}
    by_id = {item.get("suggested_id"): item for item in items}
    unknown = selected - set(by_id)
    if unknown:
        raise HarryError("The selected amendment proposal no longer exists: " + ", ".join(sorted(unknown)))
    relevant = [by_id[value] for value in selected if value in by_id]
    if not relevant:
        relevant = [item for _, _, item in sorted(scored, key=lambda row: (-row[0], row[1]))[:4]]
    if not relevant:
        relevant = items[:2]
    fields = ("kind", "name", "suggested_id", "description", "continuity_note", "positive_prompt", "negative_prompt", "beat", "reused_from", "camera_direction", "lighting_direction")
    compact = [{field: item.get(field, "") for field in fields} for item in relevant]
    index = [{"suggested_id": item.get("suggested_id", ""), "kind": item.get("kind", ""), "name": item.get("name", ""), "beat": item.get("beat", "")} for item in items]
    return compact, index


def revise_plan(provider, plan, feedback, source_excerpt="", selected_ids=None, on_progress=None):
    """Apply one grounded, scoped patch without resubmitting the call sheet."""
    feedback = (feedback or "").strip()
    if not feedback:
        raise HarryError("Give Harry The Helper a specific note before asking for a revision.")
    context = (source_excerpt or "").strip()
    relevant, proposal_index = _amendment_context(plan, feedback, context, selected_ids)
    system = f"""{HARRY_PERSONA}
You are applying one fresh, narrow amendment to an established production baseline.
The director's selected script passage is the grounding evidence. Treat it as a complete local production review: identify all characters, character states, backdrops, props, and coverage the passage needs to be filmable. Do not revisit or reinterpret unrelated story moments.

Rules:
- The baseline is already approved. Everything not returned stays exactly as it is.
- Amend any existing proposal directly required by the selected passage. Use the full call-sheet index to find its ID.
- Add a character, backdrop, prop, character state, or shot whenever the selected passage requires one that the baseline does not adequately cover.
- Use judgment: amend a continuing asset when its identity remains the same; add a distinct proposal when the passage introduces a new asset or a materially distinct filmable state.
- New SHOT IDs must preserve chronological order. For example, additions after 2.1 and before 2.2 must be named 2.1a, 2.1b, then 2.1c.
- Do not change the production summary or clarification questions.
- For every new or amended SHOT, camera_direction and lighting_direction are required.
{ERA_NOTE}

Return JSON only:
{{
  "updates": [{{"kind":"...", "suggested_id":"existing ID", "name":"...", "description":"...", "continuity_note":"...", "positive_prompt":"...", "negative_prompt":"...", "beat":"...", "reused_from":"...", "camera_direction":"for SHOT", "lighting_direction":"for SHOT"}}],
  "additions": [{{"kind":"CHARACTER|BACKDROP|PROP|SHOT", "suggested_id":"new valid ID", "name":"...", "description":"...", "continuity_note":"...", "positive_prompt":"...", "negative_prompt":"...", "beat":"...", "reused_from":"", "camera_direction":"for SHOT", "lighting_direction":"for SHOT"}}]
}}"""
    user = "DIRECTOR NOTE:\n" + feedback
    if context:
        user += "\n\nSELECTED SCRIPT PASSAGE (grounding evidence):\n" + context
    user += "\n\nRELEVANT BASELINE PROPOSALS (the only existing proposals you may amend):\n" + json.dumps(relevant, separators=(",", ":"))
    user += "\n\nCALL-SHEET ID INDEX (lookup only; do not amend these unless they are listed above):\n" + json.dumps(proposal_index, separators=(",", ":"))
    user += "\n\nReturn the scoped patch only."
    raw = ""
    try:
        if on_progress:
            on_progress("Harry The Helper is applying a focused amendment to the selected passage.")
        raw = _chat(provider, system, user)
        data = _parse_json_block(raw)
        if not isinstance(data, dict):
            raise HarryError("Harry returned an unusable amendment object.")
        updates = _clean_items(data.get("updates", []))
        additions = _clean_items(data.get("additions", []))
        allowed = {item.get("suggested_id") for item in (plan or {}).get("items", [])}
        unexpected = [item.get("suggested_id") for item in updates if item.get("suggested_id") not in allowed]
        if unexpected:
            raise HarryError("Harry tried to amend outside the selected scope: " + ", ".join(unexpected))
        revised = _merge_revision(plan, updates, additions, None, None)
    except HarryRateLimitError:
        if on_progress:
            on_progress("Amendment paused: the provider is rate limiting requests. No retry was sent.")
        raise
    except HarryError as error:
        _write_analysis_diagnostic("Director feedback", error, raw)
        raise HarryError("Harry The Helper could not apply this focused amendment. The active call sheet was left unchanged. " + str(error)) from error
    if on_progress:
        changed = len(revised.get("items", [])) - len((plan or {}).get("items", []))
        on_progress("Focused amendment complete: " + (str(changed) + " new proposal(s) added." if changed else "only the selected proposals were updated."))
    revised.update({"plan_id": datetime.now().strftime("%Y%m%d_%H%M%S"), "title": (plan or {}).get("title", "Untitled project"), "provider": provider, "era": (plan or {}).get("era", ""), "created_at": _timestamp(), "source_path": (plan or {}).get("source_path", "")})
    LIBRARY_DIR.mkdir(exist_ok=True)
    (LIBRARY_DIR / f"plan_{revised['plan_id']}.json").write_text(json.dumps(revised, indent=2), encoding="utf-8")
    return revised


def plan_to_rows(plan):
    return [[True, item["kind"], item["name"], item["suggested_id"], item["description"], item["continuity_note"], item["positive_prompt"], item["negative_prompt"], item["beat"], item["reused_from"], item.get("camera_direction", "")] for item in plan.get("items", [])]


CALL_SHEET_COLUMNS = ["Use", "Type", "Name", "Suggested ID", "Description", "Continuity note",
                      "Positive prompt", "Negative prompt", "Beat", "Chained from"]


def export_call_sheet(title, rows, plan):
    """Writes the current call sheet (as it stands in the review table --
    including any manual edits/unticks, not just the raw provider output)
    to a real .xlsx file for offline review or sharing. Two sheets:
    Call Sheet (the rows themselves) and Summary (Harry's reading + any
    clarification questions), so the export is self-contained -- someone
    reviewing it doesn't need the app open to understand the context."""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    LIBRARY_DIR.mkdir(exist_ok=True)
    export_dir = LIBRARY_DIR / "exports"
    export_dir.mkdir(exist_ok=True)

    wb = Workbook()
    sheet = wb.active
    sheet.title = "Call Sheet"
    sheet.append(CALL_SHEET_COLUMNS)
    for row in (rows or []):
        # rows may come from the Gradio Dataframe as-is; pad/trim defensively
        # in case of manual edits that changed column count somehow.
        padded = list(row) + [""] * (len(CALL_SHEET_COLUMNS) - len(row))
        sheet.append(padded[:len(CALL_SHEET_COLUMNS)])
    for i, _ in enumerate(CALL_SHEET_COLUMNS, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = 24

    summary_sheet = wb.create_sheet("Summary")
    summary_sheet.append(["Title", title or "Untitled project"])
    summary_sheet.append(["Exported", _timestamp()])
    summary_sheet.append(["Provider", (plan or {}).get("provider", "")])
    summary_sheet.append(["Era / setting", (plan or {}).get("era", "")])
    summary_sheet.append([])
    summary_sheet.append(["Harry's reading"])
    summary_sheet.append([(plan or {}).get("summary", "")])
    summary_sheet.append([])
    summary_sheet.append(["Clarification questions"])
    for question in (plan or {}).get("questions", []):
        summary_sheet.append([question])
    summary_sheet.column_dimensions["A"].width = 90

    safe_title = _slug(title) or "call_sheet"
    filename = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{safe_title}_call_sheet.xlsx"
    out_path = export_dir / filename
    wb.save(out_path)
    return str(out_path)


def rows_to_plan(rows, template=None):
    items = []
    for row in rows or []:
        if not row or len(row) < 10 or not row[0]:
            continue
        items.append(dict(zip(("selected", "kind", "name", "suggested_id", "description", "continuity_note", "positive_prompt", "negative_prompt", "beat", "reused_from"), row)))
    return {"summary": (template or {}).get("summary", ""), "questions": (template or {}).get("questions", []), "items": items}


def preset_path():
    return LIBRARY_DIR / "approved_build_presets.json"


def save_presets(plan):
    LIBRARY_DIR.mkdir(exist_ok=True)
    preset_path().write_text(json.dumps(plan, indent=2), encoding="utf-8")


def load_presets():
    if not preset_path().exists():
        return []
    return json.loads(preset_path().read_text(encoding="utf-8")).get("items", [])


def preset_choices():
    return [(f"{item.get('kind')}: {item.get('name')}", index) for index, item in enumerate(load_presets())]


def checklist_choices():
    return [f"{item.get('kind')}: {item.get('name')}" for item in load_presets()]
